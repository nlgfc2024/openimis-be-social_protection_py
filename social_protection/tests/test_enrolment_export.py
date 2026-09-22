from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook

from social_protection.export_mixin import (
    ENROLMENT_EXPORT_COLUMNS,
    ExportableSocialProtectionQueryMixin,
)


class EnrolmentExportWorkbookTest(SimpleTestCase):
    def test_workbook_contains_every_required_reporting_column(self):
        workbook_content = ExportableSocialProtectionQueryMixin._enrolment_workbook([])
        worksheet = load_workbook(BytesIO(workbook_content))["Enrolments"]
        headers = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]

        self.assertEqual(headers, list(ENROLMENT_EXPORT_COLUMNS))
        self.assertIn("date_of_birth", headers)
        self.assertIn("project_code", headers)
        self.assertIn("relationship", headers)
